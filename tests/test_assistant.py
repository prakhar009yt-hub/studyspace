import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from muj_study.core import App, StudyError, safe_name
from muj_study.gemini import Gemini, document_parts
from muj_study.lms import LMS
from muj_study.schemas import Analysis
from muj_study.youtube import YouTube, duration_seconds


class AssistantTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.app = App(self.base / 'config.json', self.base / 'data')
        self.app.cfg['request_delay'] = 0

    def tearDown(self):
        self.app.close()
        self.temp.cleanup()

    def test_cache_expiry_and_stale_fallback(self):
        self.app.cache('a', {'ok': True}, days=-1)
        self.assertIsNone(self.app.cached('a'))
        self.assertEqual(self.app.cached('a', stale=True), {'ok': True})

    def test_budget_cannot_overspend(self):
        self.app.reserve('search', 1)
        with self.assertRaises(StudyError): self.app.reserve('search', 1)
        self.assertEqual(self.app.db.execute('SELECT used FROM quota').fetchone()[0], 1)

    def test_safe_names_and_duration(self):
        self.assertEqual(safe_name('CON.pdf'), '_CON.pdf')
        self.assertNotIn('/', safe_name('../../note.pdf'))
        self.assertEqual(duration_seconds('PT1H12M30S'), 4350)
        self.assertEqual(duration_seconds('PT50S'), 50)

    def test_nested_brightspace_content(self):
        modules = [{'Title': 'Unit 1', 'Topics': [{'Title': 'A'}], 'Modules': [{'Title': 'Chapter 2', 'Topics': [{'Title': 'B'}]}]}]
        result = list(LMS.flatten_modules(modules))
        self.assertEqual(result[1][0], ('Unit 1', 'Chapter 2'))
        self.assertEqual(len(result), 2)

    def test_pdf_binary_chunking_preserves_page_counts(self):
        from pypdf import PdfReader, PdfWriter
        path = self.base / 'notes.pdf'
        writer = PdfWriter()
        for _ in range(5): writer.add_blank_page(width=200, height=200)
        writer.write(path)
        chunks = list(document_parts(path, 2))
        self.assertEqual([start for _, start in chunks], [1, 3, 5])
        self.assertEqual([len(PdfReader(io.BytesIO(part.inline_data.data)).pages) for part, _ in chunks], [2, 2, 1])

    def test_binary_pdf_api_payload_and_cached_schema_validation(self):
        from google.genai import types
        from unittest.mock import Mock
        client = Mock()
        client.__enter__ = Mock(return_value=client)
        client.__exit__ = Mock(return_value=False)
        client.models.generate_content.return_value.text = Analysis(summary='PDF summary', topics=[]).model_dump_json()
        binary = types.Part.from_bytes(data=b'%PDF\x80\xff', mime_type='application/pdf')
        with patch('google.genai.Client', return_value=client), patch.object(self.app, 'key', return_value='fake-key'):
            first = Gemini(self.app).generate('analyze', Analysis, [binary])
            second = Gemini(self.app).generate('analyze', Analysis, [binary])
        self.assertEqual(first, second)
        self.assertEqual(client.models.generate_content.call_count, 1)
        self.assertEqual(client.models.generate_content.call_args.kwargs['contents'][1].inline_data.data, b'%PDF\x80\xff')

    def test_presentation_keeps_slide_images_for_gemini(self):
        from pptx import Presentation
        from pptx.util import Inches
        import base64
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        image = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aP2kAAAAASUVORK5CYII=')
        slide.shapes.add_picture(io.BytesIO(image), Inches(1), Inches(1), width=Inches(2), height=Inches(2))
        path = self.base / 'scanned-slides.pptx'; presentation.save(path)
        parts = list(document_parts(path))
        self.assertEqual(parts[0][1], 1)
        self.assertTrue(any(p.inline_data and p.inline_data.data == image for p in parts[0][0].parts))
        self.assertIn('Slide 1', parts[0][0].parts[0].text)

    def test_analysis_cache_and_context_invalidation(self):
        path = self.base / 'notes.txt'; path.write_text('SQL joins')
        fake = Analysis(summary='Joins', topics=[])
        with patch.object(Gemini, 'generate', return_value=fake) as generate:
            first = Gemini(self.app).analyze(path, 'DBMS')
            second = Gemini(self.app).analyze(path, 'DBMS')
            self.assertEqual(first['id'], second['id'])
            self.assertTrue(second['reused'])
            self.assertEqual(generate.call_count, 1)
            self.app.cfg['level'] = 'beginner'
            third = Gemini(self.app).analyze(path, 'DBMS')
            self.assertFalse(third['reused'])
            self.assertEqual(generate.call_count, 2)
        path.write_text('changed')
        with self.assertRaises(StudyError): Gemini(self.app).note(first['id'])

    def test_feedback_applies_to_cached_recommendations(self):
        topic = {'name': 'SQL', 'queries': ['sql tutorial', 'sql examples', 'sql exam']}
        yt = YouTube(self.app)
        candidate = {'id': 'v1', 'channelId': 'c1', 'channelTitle': 'Teacher', 'score': 8, 'role': 'quick intro', 'duration_seconds': 600, 'language': 'en'}
        from muj_study.core import digest
        context = {k: self.app.cfg[k] for k in ['subject', 'level', 'languages', 'preferred_channels', 'model', 'ranking_model', 'minimum_views', 'minimum_like_ratio']}
        context['ranking_version'] = 'ai-topic-v3'
        self.app.cache('recommend:' + digest([topic, context, False]), {'topic': 'SQL', 'videos': [candidate], 'playlists': []})
        self.assertEqual(len(yt.recommend(topic)['videos']), 1)
        self.app.feedback('v1', 'c1', status='watched')
        self.assertEqual(yt.recommend(topic)['videos'], [])
        self.app.feedback('v1', status='unwatched')
        self.app.cfg['blocked_channels'] = ['c1']
        self.assertEqual(yt.recommend(topic)['videos'], [])

    def test_filter_languages_duration_and_missing_likes(self):
        v = {'id': 'v', 'snippet': {'title': 'SQL joins explained', 'channelTitle': 'Teacher', 'channelId': 'c', 'defaultAudioLanguage': 'en'}, 'statistics': {'viewCount': '2000'}, 'contentDetails': {'duration': 'PT10M'}}
        yt = YouTube(self.app)
        self.assertTrue(yt.filter(v, {}))
        v['contentDetails']['duration'] = 'PT30S'
        self.assertFalse(yt.filter(v, {}))
        v['contentDetails']['duration'] = 'PT10M'; v['snippet']['defaultAudioLanguage'] = 'fr'
        self.assertFalse(yt.filter(v, {}))

    def test_untrusted_lms_hosts_rejected(self):
        lms = LMS(self.app)
        self.assertTrue(lms.allowed('https://mujlms.manipal.edu/content/a.pdf'))
        self.assertFalse(lms.allowed('https://mujlms.manipal.edu.evil.example/a.pdf'))
        self.assertFalse(lms.allowed('file:///etc/passwd'))

    def test_download_new_then_unchanged_and_nested_paths(self):
        class Response:
            status = 200
            url = 'https://mujlms.manipal.edu/content/notes.pdf'
            headers = {'content-type': 'application/pdf', 'etag': 'abc'}
            def body(self): return b'%PDF-test'
            def dispose(self): pass
        from unittest.mock import Mock
        response = Response()
        context = Mock()
        context.request.get.return_value = response
        lms = LMS(self.app)
        self.assertEqual(lms.download(context, response.url, 'DBMS', 'Unit 1 / Chapter 2', 'Notes'), 'new')
        stored = self.app.db.execute('SELECT path FROM files').fetchone()[0]
        self.assertIn(str(Path('Unit 1') / 'Chapter 2'), stored)
        self.assertTrue(Path(stored).exists())
        response.status = 304
        self.assertEqual(lms.download(context, response.url, 'DBMS', 'Unit 1 / Chapter 2', 'Notes'), 'unchanged')
        self.assertEqual(context.request.get.call_args.kwargs['headers']['If-None-Match'], 'abc')

    def test_download_redirect_does_not_contact_untrusted_host(self):
        from unittest.mock import Mock
        context = Mock()
        response = Mock(status=302, headers={'location': 'https://evil.example/note.pdf'})
        context.request.get.return_value = response
        with self.assertRaises(StudyError):
            LMS(self.app).download(context, 'https://mujlms.manipal.edu/content/a', 'DBMS', 'Unit 1', 'A')
        self.assertEqual(context.request.get.call_count, 1)

    def test_teacher_html_page_is_saved_as_link_without_download_error(self):
        from unittest.mock import Mock
        url = 'https://mujlms.manipal.edu/content/teacher-page'
        response = Mock(status=200, url=url, headers={'content-type': 'text/html'})
        response.text.return_value = '<html><h1>Reading instructions</h1></html>'
        context = Mock()
        context.request.get.return_value = response
        self.assertEqual(LMS(self.app).download(context, url, 'DBMS', 'Unit 1', 'Reading', 3), 'linked')
        self.assertEqual(self.app.db.execute('SELECT COUNT(*) FROM links WHERE url=?', (url,)).fetchone()[0], 1)
        self.assertEqual(self.app.db.execute('SELECT COUNT(*) FROM files').fetchone()[0], 0)

    def test_missing_attachment_is_not_retried_and_reports_http_status(self):
        from unittest.mock import Mock
        context = Mock()
        context.request.get.return_value = Mock(status=404, url='https://mujlms.manipal.edu/missing', headers={})
        with self.assertRaisesRegex(StudyError, 'HTTP 404'):
            LMS(self.app).download(context, 'https://mujlms.manipal.edu/missing', 'DBMS', 'Unit', 'Missing')
        self.assertEqual(context.request.get.call_count, 1)

    def test_failed_youtube_request_uses_stale_cache_and_queues(self):
        from muj_study.core import digest
        params = {'q': 'SQL'}
        key = 'youtube:' + digest(['search', params])
        self.app.cache(key, {'items': []}, days=-1)
        with patch.object(self.app, 'key', return_value='not-real'), patch.object(self.app, 'reserve', side_effect=StudyError('Budget reached')):
            self.assertEqual(YouTube(self.app).request('search', params), {'items': []})
        self.assertEqual(self.app.db.execute('SELECT COUNT(*) FROM pending').fetchone()[0], 1)

    def test_video_candidates_survive_exhausted_gemini_budget(self):
        topic = {'name': 'SQL joins', 'queries': ['SQL joins explained', 'SQL joins examples', 'SQL joins revision']}
        video = {'id': 'match', 'snippet': {'title': 'SQL joins explained with examples', 'channelId': 'teacher', 'channelTitle': 'Teacher', 'defaultAudioLanguage': 'en', 'description': 'SQL joins for undergraduate students', 'publishedAt': '2025-01-01T00:00:00Z'}, 'statistics': {'viewCount': '10000'}, 'contentDetails': {'duration': 'PT10M', 'caption': 'true'}}
        yt = YouTube(self.app)
        with patch.object(yt, 'search', return_value=[{'id': {'videoId': 'match'}}]), patch.object(yt, 'details', return_value=[video]), patch.object(yt, 'channel_details', return_value={}), patch.object(Gemini, 'generate', side_effect=StudyError('Local gemini_calls daily budget reached')):
            result = yt.recommend(topic, quick=True)
        self.assertEqual(len(result['videos']), 1)
        self.assertEqual(result['videos'][0]['ranking_method'], 'metadata fallback')
        self.assertTrue(result['warnings'])

    def test_topic_supplements_one_match_to_three_without_losing_it(self):
        self.app.cfg['languages'] = ['en']
        topic = {'name': 'SQL joins', 'queries': ['SQL joins explained', 'SQL examples', 'SQL revision']}
        def video(vid):
            return {'id': vid, 'snippet': {'title': 'SQL joins explained with examples', 'channelId': vid, 'channelTitle': 'Teacher', 'defaultAudioLanguage': 'en', 'description': 'SQL joins', 'publishedAt': '2025-01-01T00:00:00Z'}, 'statistics': {'viewCount': '10000'}, 'contentDetails': {'duration': 'PT10M'}}
        yt = YouTube(self.app)
        with patch.object(yt, 'search', return_value=[{'id': {'videoId': v}} for v in ['a', 'b', 'c']]), patch.object(yt, 'details', side_effect=[[video('a')], [video('b'), video('c')]]), patch.object(yt, 'channel_details', return_value={}), patch.object(Gemini, 'generate', side_effect=StudyError('API service unavailable')):
            result = yt.recommend(topic, quick=True)
        self.assertEqual({v['id'] for v in result['videos']}, {'a', 'b', 'c'})

    def test_top_matches_do_not_get_displaced_by_weak_roles_or_popularity(self):
        from muj_study.core import digest
        topic = {'name': 'SQL'}
        context = {k: self.app.cfg[k] for k in ['subject', 'level', 'languages', 'preferred_channels', 'model', 'ranking_model', 'minimum_views', 'minimum_like_ratio']}
        context['ranking_version'] = 'ai-topic-v3'
        def candidate(vid, score, relevance, role='deep dive', views=1000):
            return {'id': vid, 'channelId': vid, 'channelTitle': 'Teacher', 'score': score, 'relevance': relevance, 'role': role, 'duration_seconds': 600, 'language': 'en', 'views': views}
        candidates = [candidate('best', 9, 9), candidate('second', 8.5, 9), candidate('third', 8, 8),
                      candidate('weak-role', 5, 6, 'quick intro'), candidate('viral-unrelated', 9.5, 3, 'solved problems', 100000000)]
        self.app.cache('recommend:' + digest([topic, context, False]), {'topic': 'SQL', 'videos': candidates, 'playlists': []})
        result = YouTube(self.app).recommend(topic)
        self.assertEqual([v['id'] for v in result['videos']], ['best', 'second', 'third'])

    def test_ai_plans_precise_queries_then_ranks_actual_youtube_candidates(self):
        from muj_study.schemas import VideoSearchPlan, Scores, Score
        topic = {'name': 'Carry Look-ahead Adder (CLA)', 'course': 'Computer Organization', 'concepts': ['carry propagation'], 'queries': ['complexity lecture', 'complexity example', 'revision']}
        self.app.cfg['languages'] = ['en']
        query = 'Carry lookahead adder circuit computer organization'
        def generate(prompt, schema, *args, **kwargs):
            if schema is VideoSearchPlan: return VideoSearchPlan(queries=[query, query + ' worked example'])
            return Scores(scores=[Score(video_id=v, relevance=9, clarity=8, level_match=9, depth=8, examples=8, role='deep dive', reason='Matches carry propagation', confidence='high', language='en') for v in ['a','b','c']])
        videos = [{'id': v, 'snippet': {'title': query, 'channelId': v, 'channelTitle': 'Teacher', 'defaultAudioLanguage': 'en', 'description': 'CLA circuit', 'publishedAt': '2025-01-01T00:00:00Z'}, 'statistics': {'viewCount': '10000'}, 'contentDetails': {'duration': 'PT10M'}} for v in ['a','b','c']]
        yt = YouTube(self.app)
        with patch.object(yt, 'search', return_value=[{'id': {'videoId': v}} for v in ['a','b','c']]) as search, patch.object(yt, 'details', return_value=videos), patch.object(yt, 'channel_details', return_value={}), patch.object(Gemini, 'generate', side_effect=generate):
            result = yt.recommend(topic, quick=True)
        self.assertTrue(all(query in call.args[0] for call in search.call_args_list))
        self.assertTrue(result['ai_ranked'])
        self.assertFalse(result['no_strong_match'])
        self.assertEqual(len(result['videos']), 3)

    def test_youtube_provider_mode_continues_past_local_search_and_metadata_caps(self):
        from unittest.mock import Mock
        self.app.cfg.update(youtube_use_provider_quota=True, daily_search_limit=1, daily_metadata_limit=1)
        self.app.reserve('youtube_searches', 1)
        self.app.reserve('youtube_metadata', 1)
        client = Mock(); client.__enter__ = Mock(return_value=client); client.__exit__ = Mock(return_value=False)
        response = Mock(status_code=200); response.json.return_value = {'items': []}
        client.get.return_value = response
        with patch.object(self.app, 'key', return_value='fake'), patch('httpx.Client', return_value=client):
            yt = YouTube(self.app)
            self.assertEqual(yt.request('search', {'q': 'arrays'}), {'items': []})
            self.assertEqual(yt.request('videos', {'id': 'test'}), {'items': []})
        self.assertEqual(client.get.call_count, 2)
        counts = dict(self.app.db.execute('SELECT bucket,used FROM quota'))
        self.assertEqual(counts['youtube_searches'], 2)
        self.assertEqual(counts['youtube_metadata'], 2)

    def test_provider_daily_quota_error_pauses_further_youtube_requests(self):
        from unittest.mock import Mock
        self.app.cfg['youtube_use_provider_quota'] = True
        client = Mock(); client.__enter__ = Mock(return_value=client); client.__exit__ = Mock(return_value=False)
        response = Mock(status_code=403)
        response.json.return_value = {'error': {'errors': [{'reason': 'quotaExceeded'}]}}
        client.get.return_value = response
        with patch.object(self.app, 'key', return_value='fake'), patch('httpx.Client', return_value=client):
            yt = YouTube(self.app)
            with self.assertRaisesRegex(StudyError, 'quotaExceeded'): yt.request('search', {'q': 'arrays'})
            with self.assertRaisesRegex(StudyError, 'paused'): yt.request('search', {'q': 'trees'})
        self.assertEqual(client.get.call_count, 1)

    def test_failed_topic_refresh_preserves_saved_results_and_other_topics(self):
        topics = [{'name': 'SQL joins'}, {'name': 'Normalization'}]
        saved = [{'topic': t['name'], 'videos': [{'id': str(i), 'channelId': 'c', 'channelTitle': 'Teacher'}]} for i, t in enumerate(topics)]
        self.app.db.execute('INSERT INTO history(created,note,result) VALUES(?,?,?)', (1, 'note', json.dumps(saved)))
        self.app.db.commit()
        yt = YouTube(self.app)
        with patch.object(Gemini, 'note', return_value={'analysis': json.dumps({'topics': topics})}), patch.object(yt, 'recommend', return_value={'topic': 'SQL joins', 'videos': [], 'warnings': ['Search budget reached']}):
            result = yt.note_topic('note', 0, refresh=True)
        self.assertEqual(len(result['videos']), 1)
        latest = json.loads(self.app.db.execute('SELECT result FROM history ORDER BY id DESC LIMIT 1').fetchone()['result'])
        self.assertEqual(len(latest), 2)


if __name__ == '__main__':
    unittest.main()
